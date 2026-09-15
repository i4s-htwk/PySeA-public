from dataclasses import dataclass
from typing import List, Optional
from openpyxl import load_workbook
from openpyxl.utils import coordinate_to_tuple


@dataclass
class ExcelVariable:
    """A single variant-dependent value read from an Excel file.

    Instances are held in ``ExcelVariables.variables`` and created via
    ``ExcelVariables.add_variable()``. Each variable reads its values
    horizontally starting at ``cell``, with one column per variant. It can
    be embedded directly into an item body or table cell.
    """
    variable_id: str
    cell: Optional[str] = None

    def to_dict(self) -> dict:
        return {
            "variable_id": self.variable_id,
            "cell": self.cell
        }

    @staticmethod
    def from_dict(data):
        if data is None:
            return None

        if isinstance(data, ExcelVariable):
            return data

        if isinstance(data, (list, tuple)):
            variable_id = data[0] if len(data) > 0 else None
            cell = data[1] if len(data) > 1 else None
            return ExcelVariable(variable_id=variable_id, cell=cell)

        return ExcelVariable(
            variable_id=data.get("variable_id"),
            cell=data.get("cell")
        )

    def __str__(self) -> str:
        return "{" + self.variable_id + "}"

    def __radd__(self, other: object) -> str:
        if isinstance(other, str):
            return other + str(self)
        return NotImplemented

    def __add__(self, other: object) -> str:
        if isinstance(other, str):
            return str(self) + other
        return NotImplemented

@dataclass
class ExcelVariables:
    """Manages the Excel-based variables of a section or task.

    An instance of this class is returned by
    ``SectionFacade.add_excel_variables()`` / ``TaskFacade.add_excel_variables()``.
    ``excel_file`` and ``page`` select the data source, and
    ``add_variable()`` creates an ``ExcelVariable`` for a given start cell.
    """
    item_nr: int = None
    excel_file: str = None
    page: str = None
    variables: List[ExcelVariable] = None

    def to_dict(self) -> dict:
        return {
            "item_nr": self.item_nr,
            "excel_file": self.excel_file,
            "page": self.page,
            "variables": [variable.to_dict() for variable in self.variables] if self.variables else []
        }

    @staticmethod
    def from_dict(data):
        if data is None:
            return None

        ev = ExcelVariables(item_nr=data.get("item_nr"))
        ev.excel_file = data.get("excel_file")
        ev.page = data.get("page")
        variables = data.get("variables", [])
        ev.variables = [ExcelVariable.from_dict(variable) for variable in variables] if variables else []

        return ev

    def get_nks(self, variants_nr: int, excel_cell: str):
        wb = load_workbook(self.excel_file, data_only=False)
        ws = wb[self.page]
        start_row, start_col = coordinate_to_tuple(excel_cell)

        most_nks = 0
        for i in range(variants_nr):
            cell = ws.cell(row=start_row, column=start_col + i)
            value = str(cell.value)
            if "." in value:
                nks = len(value.split(".", 1)[1])
                if nks > most_nks:
                    most_nks = nks
            elif "," in value:
                nks = len(value.split(",", 1)[1])
                if nks > most_nks:
                    most_nks = nks

        wb.close()
        return most_nks

    def maxima_code(self, variants_nr: int, excel_cell: str) -> str:
        wb = load_workbook(self.excel_file, data_only=False)
        ws = wb[self.page]
        start_row, start_col = coordinate_to_tuple(excel_cell)

        values = []
        for i in range(int(variants_nr)):
            cell = ws.cell(row=start_row, column=start_col + i)
            value = cell.value
            values.append(value)

        wb.close()

        parts = []
        for idx, value in enumerate(values, start=1):
            if idx == 1:
                parts.append(f'if $(1) = {idx} then {value}')
            else:
                parts.append(f'else if $(1) = {idx} then {value}')

        maxima_str = f'float({" ".join(parts)});'
        return maxima_str

    def add_variable(self, cell=None):
        """Add a new Excel-based variable.

        The variable reads its values horizontally starting at ``cell``:
        that cell holds the value for variant 1, the cell to its right the
        value for variant 2, and so on.
        Args:
            cell: The Excel cell containing the value for variant 1 (e.g.
                ``"C10"``).
        """
        if self.variables is None:
            self.variables = []

        index = len(self.variables) + 1
        variable_id = f"EXCEL_VARIABLE_{self.item_nr}_{index}"
        self.variables.append(ExcelVariable(variable_id=variable_id, cell=cell))

    def delete_variable(self, variable_id: str):
        if self.variables is None:
            self.variables = []
            return

        self.variables = [variable for variable in self.variables if variable.variable_id != variable_id]
        for i, variable in enumerate(self.variables, start=1):
            variable.variable_id = f"EXCEL_VARIABLE_{self.item_nr}_{i}"

    def set_variable(self, variable_id: str, cell: str):
        if self.variables is None:
            return

        for variable in self.variables:
            if variable.variable_id == variable_id:
                variable.cell = cell
                break
