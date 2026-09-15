from dataclasses import dataclass, field
from typing import List, Dict, Optional
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.workbook.workbook import Workbook

from backend.utils import Response, ExcelResponse

@dataclass
class ExcelResponses:
    """Manages the Excel-based response fields of a task.

    An instance of this class is returned by ``TaskFacade.excel_responses()``.
    ``excel_file`` and ``page`` select the data source, and
    ``add_response()`` creates a response field whose correct values per
    variant are read from a row starting at the given cell; the created
    ``ExcelResponse`` objects are then available in ``responses``.
    """
    excel_file: str
    page: str
    answer_acc_from_excel: bool

    responses: List[ExcelResponse]

    wb: Optional[Workbook] = field(default=None, init=False, repr=False)
    ws: Optional[Worksheet] = field(default=None, init=False, repr=False)

    def open(self):
        # Wenn du nur liest, hilft oft read_only=True (schneller/ weniger RAM)
        self.wb = load_workbook(self.excel_file, data_only=False)
        self.ws = self.wb[self.page]

    def close(self):
        if self.wb is not None:
            self.wb.close()
            self.wb = None
            self.ws = None

    def to_dict(self) -> dict:
        return {
            "excel_file": self.excel_file,
            "page": self.page,
            "answer_acc_from_excel": self.answer_acc_from_excel,
            "responses": [resp.to_dict() for resp in self.responses]
        }

    @staticmethod
    def from_dict(data: dict) -> "ExcelResponses":
        responses_data = data.get("responses", [])
        responses = [ExcelResponse.from_dict(rd) for rd in responses_data]
        return ExcelResponses(
            excel_file=data.get("excel_file", ""),
            page=data.get("page", ""),
            answer_acc_from_excel=data.get("answer_acc_from_excel", True),
            responses=responses
        )

    def maxima_code(self, variants_nr: int, excel_cell: str, offset:int = 0) -> str:
        '''wb = load_workbook(self.excel_file, data_only = False)
        ws = wb[self.page]'''
        start_row, start_col = coordinate_to_tuple(excel_cell)
        start_row += offset
        values = []
        for i in range(int(variants_nr)):
            cell = self.ws.cell(row=start_row, column=start_col + i)
            value = cell.value
            values.append(value)

        #wb.close()

        parts = []
        for idx, value in enumerate(values, start=1):
            if idx == 1:
                parts.append(f'if $(1) = {idx} then {value}')
            else:
                parts.append(f'else if $(1) = {idx} then {value}')

        maxima_str = f'float({" ".join(parts)});'
        return maxima_str

    def add_response(self, cell = None):
        """Add a response field whose correct values are read from Excel.

        Like ``ExcelVariables``, the response reads its values
        horizontally: ``cell`` holds the correct value for variant 1, the
        cell to its right the value for variant 2, and so on. The created
        ``ExcelResponse`` is appended to ``responses``.
        Args:
            cell: The Excel cell containing the correct value for variant 1
                (e.g. ``"C3"``).
        """
        index = len(self.responses) + 1
        response_id = f"EXCEL_RESPONSE_{index}"
        self.responses.append(ExcelResponse(id=response_id, cell=cell if cell is not None else ""))

    def delete_response(self, response_id: str):
        self.responses = [resp for resp in self.responses if resp.id != response_id]
        for i, resp in enumerate(self.responses, start=1):
            resp.id = f"EXCEL_RESPONSE_{i}"

    def set_response_value(self, response_id: str, cell: str):
        for resp in self.responses:
            if resp.id == response_id:
                resp.cell = cell
                break
    def set_response_points(self, response_id: str, points: str):
        for resp in self.responses:
            if resp.id == response_id:
                resp.points = points
                break

@dataclass
class StoreResponsesItem:
    excel_responses: Optional[ExcelResponses] = None
    responses: List[Response] = field(default_factory=list)

    def add_excel_responses(self):
        excel_responses = ExcelResponses("", "", True, [])
        self.excel_responses = excel_responses
        return excel_responses

    def add_response(self, value = None):
        index = len(self.responses) + 1
        response_id = f"RESPONSE_{index}"
        response = Response(id=response_id, value=str(value) if value is not None else "")
        self.responses.append(response)
        return response

    def delete_response(self, value_id: str):
        self.responses = [rv for rv in self.responses if rv.id != value_id]
        for i, rv in enumerate(self.responses, start=1):
            rv.id = f"RESPONSE_{i}"

    def set_response_value(self, value_id: str, value):
        for rv in self.responses:
            if rv.id == value_id:
                if value in (None, ""):
                    rv.value = ""
                else:
                    try:
                        rv.value = str(value)
                    except (ValueError, TypeError):
                        rv.value = ""
                break

    def get_merged_responses(self, selection, matching, graphical_assignment):
        merged_responses = []
        for response in self.responses:
            merged_responses.append(response)
        if self.excel_responses:
            for response in self.excel_responses.responses:
                merged_responses.append(response)
        if selection:
            merged_responses.append(selection)
        if matching:
            merged_responses.append(matching)
        if graphical_assignment:
            merged_responses.append(graphical_assignment)
        return merged_responses

@dataclass
class StoreResponses:
    items: Dict[str, StoreResponsesItem] = field(default_factory=dict)

    def add_item(self, item_id: str):
        self.items[item_id] = StoreResponsesItem()

    def delete_item(self, id_map: dict[str, str]):
        new_items = {}
        for old_id, store_item in self.items.items():
            new_id = id_map.get(old_id, old_id)
            if new_id is None:
                continue
            new_items[new_id] = store_item
        self.items = new_items

    def get_item(self, item_id: str) -> StoreResponsesItem:
        return self.items.get(item_id)

    def get_responses(self, item_id: str) -> List[Response]:
        store_item = self.items.get(item_id)
        if store_item:
            return store_item.responses
        return []

    def get_excel_responses(self, item_id: str) -> ExcelResponses:
        store_item = self.items.get(item_id)
        if store_item:
            return store_item.excel_responses
        return None

    def to_dict(self) -> Dict[str, dict]:
        return {
            item_id: {
                "excel_responses": store_item.excel_responses.to_dict() if store_item.excel_responses else None,
                "responses": [rv.to_dict() for rv in store_item.responses]
            }
            for item_id, store_item in self.items.items()
        }

    @classmethod
    def from_dict(cls, data: Dict[str, dict]) -> "StoreResponses":
        items = {}
        for item_id, item_data in data.items():
            excel_responses_data = item_data.get("excel_responses")
            excel_responses = (ExcelResponses.from_dict(excel_responses_data) if excel_responses_data else None)
            responses_data = item_data.get("responses", [])
            responses = [Response.from_dict(rd) for rd in responses_data]

            items[item_id] = StoreResponsesItem(excel_responses=excel_responses, responses=responses)
        return cls(items=items)

    def open_excelfiles(self):
        for id, item in self.items.items():
            if item.excel_responses:
                item.excel_responses.open()

    def close_excelfiles(self):
        for id, item in self.items.items():
            if item.excel_responses:
                item.excel_responses.close()

    def __getitem__(self, item_id: str) -> StoreResponsesItem:
        return self.items[item_id]

    def __str__(self):
        return str({
            eid: {
                "excel_responses": excel_responses.excel_responses,
                "responses": [rv for rv in excel_responses.responses]
            }
            for eid, excel_responses in self.items.items()
        })