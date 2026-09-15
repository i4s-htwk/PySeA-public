################################
from openpyxl import Workbook
from openpyxl.styles import numbers

##############################################################
# Erweiterung für Aufgabe 3 - Tabelle aus Excel-Datei erzeugen
# Sie können diesen Codeblock in Ihre Aufgabe 3 einfügen, um die Excel-Datei "produkte_tabelle.xlsx" zu erstellen, die für die Aufgabe benötigt wird.
##############################################################

###############################
# Excel-Datei erstellen
###############################

# Variablen für Spalten und Zeilen
A, B, C = 13.0, 27.0, 41.0
D, E, F = 0.2, 0.7, 0.3

# Excel-Datei erstellen
wb = Workbook()
ws = wb.active

# Blatt umbenennen
ws.title = "Produkte"

# Spalten- und Zeilenüberschriften schreiben
ws['B1'] = 'A'
ws['C1'] = 'B'
ws['D1'] = 'C'
ws['A2'] = 'D'
ws['A3'] = 'E'
ws['A4'] = 'F'
ws['A5'] = 'Variable A = '
ws['A6'] = 'Variable B = '
ws['A7'] = 'Variable C = '
ws['A8'] = 'Variable D = '
ws['A9'] = 'Variable E = '
ws['A10'] = 'Variable F = '

# Variablen in die Tabelle schreiben
ws['B5'] = A
ws['B6'] = B
ws['B7'] = C
ws['B8'] = D
ws['B9'] = E
ws['B10'] = F

# Produkte berechnen und in die Tabelle schreiben
ws['B2'] = A * D
ws['C2'] = B * D
ws['D2'] = C * D

ws['B3'] = A * E
ws['C3'] = B * E
ws['D3'] = C * E

ws['B4'] = A * F
ws['C4'] = B * F
ws['D4'] = C * F

# Formatierung für Nachkommastellen anwenden
for row in ws.iter_rows(min_row=2, max_row=4, min_col=2, max_col=4):
    for cell in row:
        cell.number_format = numbers.FORMAT_NUMBER_00

# Excel-Datei speichern
excel_file = "examples/PySeA_workshop/files/produkte_tabelle.xlsx"
wb.save(excel_file)
print("Excel-Datei wurde erstellt: produkte_tabelle.xlsx")
################################