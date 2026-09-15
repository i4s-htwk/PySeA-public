import os

from dash import Output, Input, State, callback, dcc, html, ALL, ctx
from dash.exceptions import PreventUpdate
from openpyxl import load_workbook

from backend.TestStructure import TestStructure
from backend.stores import StoreConfigurations
from backend.utils import ExcelVariables


@callback(
    Output("div_variables", "children"),
    Output("view_edit_item", "data", allow_duplicate=True),

    Input("b_excel_variables", "n_clicks"),
    State("store_teststructure", "data"),
    State("store_current_obj", "data"),
    State("store_configurations", "data"),
    prevent_initial_call=True
)
def render_excel_variables(n_clicks, store_teststructure, current_obj, store_configurations):
    if not n_clicks:
        raise PreventUpdate

    test_structure = TestStructure.from_dict(store_teststructure)
    current_obj = test_structure.get_object(current_obj)
    configurations = StoreConfigurations.from_dict(store_configurations or {})
    path_excel_file = configurations.path_excel_files

    if not current_obj.excel_variables:
        result = html.Div([html.Button("Variablen aus Exceldatei verwenden", id="b_initialize_excel_variables"),
                         html.Div([dcc.Dropdown(id="excel_variable_doc", style={"display":"none"}),
                                   dcc.Dropdown(id="excel_variable_page", style={"display":"none"}),
                                   html.Button(id="b_new_excel_variable", style={"display":"none"})])],
                          style={"marginLeft": "260px", "marginTop": "60px", "paddingLeft": "20px"})
    else:
        variables = current_obj.excel_variables
        ev_div = []
        for i, ev in enumerate(variables.variables):
            ev_div.append(html.Div([ev.variable_id, dcc.Input(id={"type":"i_excel_variable_cell", "index":i}, value=ev.cell)]))

        options_excel = options_page = []
        if path_excel_file:
            options_excel = [{"label": path.split("/")[-1], "value": path} for path in path_excel_file if path]
        if variables.page:
            options_page = [{"label": variables.page, "value": variables.page}]

        result = html.Div([
            html.H4("Excelvariablen", style={"marginLeft": "10px", "marginBottom": "10px"}),
            html.Button(id="b_initialize_excel_variables", style={"display":"none"}),
                html.Div([
                    html.Div([
                        dcc.Dropdown(id="excel_variable_doc", options=options_excel,
                                     value=variables.excel_file, placeholder="Exceldatei auswählen",
                                     style={"paddingLeft": "100px", "width": "200px", "height": "15px", "fontSize": "12px"}),
                        dcc.Dropdown(id="excel_variable_page", options=options_page,
                                     value=variables.page, placeholder="Seite auswählen",
                                     style={"paddingLeft": "20px", "width": "200px", "height": "15px", "fontSize": "12px"}),
                        ],
                    style={"display": "flex", "flex-direction": "row", "marginLeft": "20", "marginTop": "15px", "marginBottom":"15px"}),
                    html.Div(ev_div, style={"display":"flex", "flex-direction":"column", "marginTop":"10px"}),
                    html.Button("neue Variable", id="b_new_excel_variable")],style={"display": "flex", "flex-direction": "column"})],
            style={"display": "flex", "alignItems": "flex-start","marginLeft": "260px", "marginTop": "60px", "paddingLeft": "20px"})
    return [result, "variables"]

@callback(
    Output("excel_variable_page", "options"),

    Input("excel_variable_doc", "value"),
    Input("variables_excel_is_saved", "data"),
    prevent_initial_call=True
)
def lade_seitenoptionen(excel_path_file, is_saved):
    if not excel_path_file or not os.path.exists(excel_path_file):
        raise PreventUpdate
    try:
        wb = load_workbook(excel_path_file, read_only=True)
        seiten = wb.sheetnames
        wb.close()
        options = [{"label": name, "value": name} for name in seiten]
        return options
    except Exception as e:
        print("Fehler beim Einlesen der Excel-Datei:", e)
        return []

@callback(
    Output("variables_excel_is_saved", "data", allow_duplicate=True),
    Output("div_variables", "children", allow_duplicate=True),

    Input("variables_excel_is_saved", "data"),
    State("store_teststructure", "data"),
    State("store_current_obj", "data"),
    State("store_configurations", "data"),
    prevent_initial_call=True
)
def rerender_excel_variables(is_saved,  store_teststructure, current_obj, store_configurations):
    if not ctx.triggered_id or not is_saved:
        raise PreventUpdate
    layout, *_ = render_excel_variables(1, store_teststructure, current_obj, store_configurations)
    return False, layout

@callback(
    Output("store_teststructure", "data", allow_duplicate=True),
    Output("variables_excel_is_saved", "data", allow_duplicate= True),

    Input("b_initialize_excel_variables", "n_clicks"),
    Input("b_new_excel_variable", "n_clicks"),
    Input("excel_variable_doc", "value"),
    Input("excel_variable_page", "value"),
    Input({"type":"i_excel_variable_cell", "index":ALL}, "value"),

    State("store_teststructure", "data"),
    State("store_current_obj", "data"),
    prevent_initial_call=True
)
def save_variables_excel(b_use_excel_variables, b_new_excel_variable, excel_variable_doc, excel_variable_page, excel_variable_cells, test_structure, current_obj):
    if not ctx.triggered_id:
        raise PreventUpdate

    test_structure = TestStructure.from_dict(test_structure)
    current_obj_ts = test_structure.get_object(current_obj)

    if ctx.triggered_id == "b_new_excel_variable":
        current_obj_ts.excel_variables.add_variable(current_obj_ts.section_nr)
    elif ctx.triggered_id == "excel_variable_doc":
        current_obj_ts.excel_variables.excel_file = excel_variable_doc
    elif ctx.triggered_id == "excel_variable_page":
        current_obj_ts.excel_variables.page = excel_variable_page
    elif isinstance(ctx.triggered_id, dict) and ctx.triggered_id["type"] == "i_excel_variable_cell":
        i = ctx.triggered_id.get("index")
        current_obj_ts.excel_variables.variables[i].cell = excel_variable_cells[i]
    elif ctx.triggered_id == "b_initialize_excel_variables":
        if not b_use_excel_variables:  # None oder 0
            raise PreventUpdate
        current_obj_ts.add_excel_variables()

    return test_structure.to_dict(), True