import os

from dash import Output, Input, State, callback, dcc, html, ALL, ctx
from dash.exceptions import PreventUpdate
from openpyxl import load_workbook

from backend.TestStructure import TestStructure
from backend.Section import Section
from backend.stores import  StoreItemBody, StoreConfigurations
from backend.utils import VariantDependentImage, VariantDependentTable


@callback(
    Output("div_general", "children"),
    Output("view_edit_item", "data", allow_duplicate=True),

    Input("b_general", "n_clicks"),
    State("store_item_body", "data"),
    State("store_teststructure", "data"),
    State("store_configurations", "data"),
    State("store_current_obj", "data"),
    prevent_initial_call=True
)
def render_general(n_clicks, item_body, teststructure, configuration, current_obj):

    item_body = StoreItemBody(items=item_body or {})
    test_structure = TestStructure.from_dict(teststructure)
    configuration = StoreConfigurations.from_dict(configuration or {})
    current_obj_ts = test_structure.get_object(current_obj)

    element = ""
    point_distribution = html.Div([dcc.Input(id="i_points_per_gap_tasks",style={"display":"none"}), dcc.Input(id="i_points_task",style={"display":"none"})])
    #excel_variables = html.Div([dcc.Dropdown(id="excel_variable_doc", style={"display":"none"}), dcc.Dropdown(id="excel_variable_page", style={"display":"none"}), html.Button(id="b_use_excel_variables", style={"display":"none"}), html.Button(id="b_new_excel_variable", style={"display":"none"})])
    variant_dependent_content = html.Div([dcc.Dropdown(id="dd_variant_dependent_content_type", style={"display":"none"}), html.Button(id="b_new_variant_dependent_content", style={"display":"none"})])
    if "section" in current_obj:
        element = "Sektion"
        #excel_variables = render_excel_variables(current_obj_ts, configuration.path_excel_files)
    elif "task" in current_obj:
        point_distribution = render_point_distribution(current_obj_ts.point_distribution)
        variant_dependent_content = render_variant_dependent_content(current_obj_ts)
        element = "Aufgabe"

    return html.Div([
        html.Div([
            html.H3("Titel der "+element, style={"whiteSpace": "nowrap", "marginRight": "15px", "marginBottom": "0"}),
            dcc.Input(id="i_title", type="text", value=current_obj_ts.title, placeholder="Titel eingeben", style={"width": "300px", "height": "35px", "marginTop": "6px"})
        ], style={"display": "flex", "alignItems": "flex-start", "marginBottom": "20px"}),
        html.Div(id="item_body_task", children=[
            html.H4("Fragenkörper"),
            dcc.Textarea(id='i_text', value=item_body.get_item(current_obj) or "",
                         style={'width': '100%', 'height': 300, 'resize': 'vertical', 'fontSize': 16})],
                 style={"marginBottom": "20px", "width": "100%"}),
        point_distribution,
        #excel_variables,
        variant_dependent_content
    ], style={"marginLeft": "260px","marginTop" :"50px","paddingLeft": "20px"}), "general"

def render_point_distribution(point_distribution):
    return html.Div([
        html.H4("Punkteverteilung", style={"marginLeft": "10px", "marginBottom": "10px"}),
        html.Div([
            html.Div([
                "Punkte auf Aufgabe:", dcc.Input(id="i_points_task", value=point_distribution["task"], style={"marginLeft": "60px", "width": "80px"})], style={"paddingBottom": "10px"}),
            html.Div([
                "Punkte pro Lücke:", dcc.Input(id="i_points_per_gap_tasks", value=point_distribution["gap"], style={"marginLeft":"72px", "width":"80px"})], style={"paddingBottom":"10px"})],
            style={"display": "flex", "flex-direction": "column", "marginLeft": "87px" ,"marginTop":"15px"})],
        style={"display": "flex", "alignItems": "flex-start", "border": "0.5px solid black", "padding": "5px",
              "marginTop": "15px"})

'''def render_excel_variables(current_obj, path_excel_file):
    if not current_obj.excel_variables:
        return html.Div([html.Button("Variablen aus Exceldatei verwenden", id="b_use_excel_variables"),
                         html.Div([dcc.Dropdown(id="excel_variable_doc", style={"display":"none"}), dcc.Dropdown(id="excel_variable_page", style={"display":"none"}), html.Button(id="b_use_excel_variables", style={"display":"none"}), html.Button(id="b_new_excel_variable", style={"display":"none"})])])

    variables = current_obj.excel_variables
    ev_div = []
    for i, ev in enumerate(variables.variables):
        ev_div.append(html.Div([ev[0], dcc.Input(id={"type":"i_excel_variable_cell", "index":i}, value=ev[1])]))

    options_excel = options_page = []
    if path_excel_file:
        options_excel = [{"label": path.split("/")[-1], "value": path} for path in path_excel_file if path]
    if variables.page:
        options_page = [{"label": variables.page, "value": variables.page}]

    return html.Div([
        html.H4("Excelvariablen", style={"marginLeft": "10px", "marginBottom": "10px"}),
        html.Button(id="b_use_excel_variables", style={"display":"none"}),
            html.Div([
                html.Div([
                    dcc.Dropdown(id="excel_variable_doc", options=options_excel,
                                 value=variables.excel_file, placeholder="Exceldatei auswählen",
                                 style={"paddingLeft": "100px", "width": "200px", "height": "15px", "fontSize": "12px"}),
                    dcc.Dropdown(id="excel_variable_page", options=options_page,
                                 value=variables.page, placeholder="Seite auswählen",
                                 style={"paddingLeft": "20px", "width": "200px", "height": "15px", "fontSize": "12px"}),
                    ],
                style={"display": "flex", "flex-direction": "row", "marginLeft": "20", "marginTop": "15px", "marginBottom":"15px"}),
                html.Div(ev_div, style={"display":"flex", "flex-direction":"column", "marginTop":"10px"}),
                html.Button("neue Variable", id="b_new_excel_variable")],style={"display": "flex", "flex-direction": "column"})],
        style={"display": "flex", "alignItems": "flex-start", "border": "0.5px solid black", "padding": "5px", "marginTop": "15px"})
'''
'''@callback(
    Output("excel_variable_page", "options"),

    Input("excel_variable_doc", "value"),
    Input("general_is_saved", "data"),
    prevent_initial_call=True
)
def lade_seitenoptionen(excel_path_file, is_saved):
    if not excel_path_file or not os.path.exists(excel_path_file):
        raise PreventUpdate
    try:
        wb = load_workbook(excel_path_file, read_only=True)
        seiten = wb.sheetnames
        wb.close()
        options = [{"label": name, "value": name} for name in seiten]
        return options
    except Exception as e:
        print("Fehler beim Einlesen der Excel-Datei:", e)
        return []'''

def render_variant_dependent_content(current_obj):
    if isinstance(current_obj, Section):
        options = ["Bild"]
        vdt_div = html.Div(style={"display":"none"})
    else:
        options = ["Bild", "Tabelle"]
        vdt_div = render_variant_dependent_tables(current_obj.variant_dependent_tables)
    return html.Div([
        html.Div([
        html.H4("Variantenabhängiger Inhalt", style={"marginLeft": "10px", "marginBottom": "10px"}),
        html.Div([
        render_variant_dependent_images(current_obj.variant_dependent_images),
        vdt_div], style={"display":"flex", "flex-direction":"column"})], style={"display":"flex", "flex-direction":"row"}),
        html.Div([
            dcc.Dropdown(id="dd_variant_dependent_content_type", options=options, style={"width": "200px", "height": "15px", "fontSize": "12px", "marginBottom":"20px"}),
            html.Button("neuer Inhalt", id="b_new_variant_dependent_content",style={"width":"100px", "marginLeft":"20px"})],style={"display":"flex", "flex-direction":"row", "align-items":"center"} )],
        style={"display": "flex", "flex-direction":"column", "border": "0.5px solid black", "padding": "5px",
              "marginTop": "15px"})

def render_variant_dependent_images(vd_images):
    div = []
    for i, vdi in enumerate(vd_images):
        div_images = []
        for j, (img_id, value) in enumerate(vdi.images.items()):
            div_images.append(html.Div([
                    dcc.Input(id={"type": "vdi_image_id", "vdi": i, "image": img_id}, value=value["image_id"], style={"width":"200px", "marginRight":"10px"}),
                    dcc.Input(id={"type": "vdi_image_lb", "vdi": i, "image": img_id, "img_idx":j}, value=value["lower"], style={"width":"30px", "marginRight":"10px"}),
                    dcc.Input(id={"type": "vdi_image_ub", "vdi": i, "image": img_id, "img_idx":j}, value=value["upper"], style={"width":"30px"}),
                ]))
        div_images.append(html.Button("neuer Wert", id={"type":"b_new_image_vdi", "index":i}, style={"width":"90px"}))
        div.append(html.Div([
            html.Div([vdi.id], style={"marginRight":"27px"}),
            html.Div(div_images, style={"display":"flex", "flex-direction":"column", "marginLeft":"20px"})
        ], style={"display":"flex", "flex-direction":"row"}))
    if len(vd_images) != 0: topic = html.H4("Variantenabhängige Bilder:", style={"margin":0})
    else: topic = None
    return html.Div(
        [topic,
         html.Div(div, style={"marginLeft":"32px"})
        ], style={"marginLeft":"20px", "display":"flex", "flex-direction":"row", "align-items":"top", "marginTop":"17px"})

def render_variant_dependent_tables(vd_tables):
    div = []
    for i, vdt in enumerate(vd_tables):
        div_tables = []
        for j, (table_id, value) in enumerate(vdt.tables.items()):
            div_tables.append(html.Div([
                    dcc.Input(id={"type": "vdt_table_id", "vdt": i, "table": table_id}, value=value["table_id"], style={"width":"200px", "marginRight":"10px"}),
                    dcc.Input(id={"type": "vdt_table_lb", "vdt": i, "table": table_id, "table_idx":j}, value=value["lower"], style={"width":"30px", "marginRight":"10px"}),
                    dcc.Input(id={"type": "vdt_table_ub", "vdt": i, "table": table_id, "table_idx":j}, value=value["upper"], style={"width":"30px"}),
                ]))
        div_tables.append(html.Button("neuer Wert", id={"type":"b_new_table_vdt", "index":i}, style={"width":"90px"}))
        div.append(html.Div([
            vdt.id,
            html.Div(div_tables, style={"display":"flex", "flex-direction":"column", "marginLeft":"20px"})
        ], style={"display":"flex", "flex-direction":"row"}))
    if len(vd_tables) != 0: topic = html.H4("Variantenabhängige Tabellen:", style={"margin":0})
    else: topic = None
    return html.Div(
        [topic,
         html.Div(div, style={"marginLeft":"20px"})
        ], style={"marginLeft":"20px", "display":"flex", "flex-direction":"row", "align-items":"top"})

@callback(
    Output("general_is_saved", "data", allow_duplicate=True),
    Output("div_general", "children", allow_duplicate=True),

    Input("general_is_saved", "data"),
    State("store_item_body", "data"),
    State("store_teststructure", "data"),
    State("store_configurations", "data"),
    State("store_current_obj", "data"),
    prevent_initial_call=True
)
def rerender_general(is_saved, store_item_body, store_teststructure, store_configurations, current_obj):
    if not ctx.triggered_id or not is_saved:
        raise PreventUpdate
    layout, *_ = render_general(1, store_item_body, store_teststructure, store_configurations, current_obj)
    return False, layout

@callback(
    Output("store_item_body", "data", allow_duplicate=True),
    Output("store_teststructure", "data", allow_duplicate=True),
    Output("general_is_saved", "data", allow_duplicate= True),

    Input("i_title", "value"),
    Input("i_text", "value"),

    Input("i_points_task", "value"),
    Input("i_points_per_gap_tasks", "value"),

    State("dd_variant_dependent_content_type", "value"),
    Input("b_new_variant_dependent_content", "n_clicks"),
    Input({"type":"b_new_image_vdi", "index":ALL}, "n_clicks"),
    Input({"type":"vdi_image_id", "vdi":ALL, "image":ALL}, "value"),
    Input({"type":"vdi_image_lb", "vdi":ALL, "image":ALL, "img_idx":ALL}, "value"),
    Input({"type":"vdi_image_ub", "vdi":ALL, "image":ALL, "img_idx":ALL}, "value"),
    Input({"type":"b_new_table_vdt", "index":ALL}, "n_clicks"),
    Input({"type":"vdt_table_id", "vdt":ALL, "table":ALL}, "value"),
    Input({"type":"vdt_table_lb", "vdt":ALL, "table":ALL, "table_idx":ALL}, "value"),
    Input({"type":"vdt_table_ub", "vdt":ALL, "table":ALL, "table_idx":ALL}, "value"),

    State("store_item_body", "data"),
    State("store_teststructure", "data"),
    State("store_current_obj", "data"),
    prevent_initial_call=True
)
def save_general(title, text, points_task, points_per_gap, variant_dependent_content_type, new_variant_dependent_content, new_image_vdi, vdi_image_id, vdi_image_lb, vdi_image_ub,
                 new_table_vdt, vdt_table_id, vdt_table_lb, vdt_table_ub,
                 store_item_body, test_structure, current_obj):
    if not current_obj:
        raise PreventUpdate

    test_structure = TestStructure.from_dict(test_structure)
    current_obj_ts = test_structure.get_object(current_obj)
    item_body = StoreItemBody(items=store_item_body or {})

    if ctx.triggered_id == "i_points_per_gap_tasks":
        if points_per_gap =="": points_per_gap = None
        current_obj_ts.point_distribution["gap"] = points_per_gap
    elif ctx.triggered_id == "i_points_task":
        if points_task =="": points_task = None
        current_obj_ts.point_distribution["task"] = points_task
    elif ctx.triggered_id == "b_new_variant_dependent_content":
        type = variant_dependent_content_type
        if type == "Bild":
            current_obj_ts.add_variant_dependent_image()
        elif type == "Tabelle":
            current_obj_ts.add_variant_dependent_table()
    elif isinstance(ctx.triggered_id, dict) and ctx.triggered_id.get("type") == "b_new_image_vdi":
        index = ctx.triggered_id.get("index")
        vdi = current_obj_ts.variant_dependent_images[index]
        vdi.add_image()
    elif isinstance(ctx.triggered_id, dict) and ctx.triggered_id.get("type") == "b_new_table_vdt":
        index = ctx.triggered_id.get("index")
        vdt = current_obj_ts.variant_dependent_tables[index]
        vdt.add_table()

    elif isinstance(ctx.triggered_id, dict) and ctx.triggered_id.get("type") == "vdi_image_id":
        index = ctx.triggered_id.get("vdi")
        vdi = current_obj_ts.variant_dependent_images[index]
        value_key = ctx.triggered_id.get("image")
        new_image_id = ctx.triggered[0]["value"]
        vdi.set_image_id(value_key, new_image_id)
    elif isinstance(ctx.triggered_id, dict) and ctx.triggered_id.get("type") == "vdt_table_id":
        index = ctx.triggered_id.get("vdt")
        vdt = current_obj_ts.variant_dependent_tables[index]
        value_key = ctx.triggered_id.get("table")
        new_table_id = ctx.triggered[0]["value"]
        vdt.set_table_id(value_key, new_table_id)

    elif isinstance(ctx.triggered_id, dict) and ctx.triggered_id.get("type") in ("vdi_image_lb", "vdi_image_ub"):
        index = ctx.triggered_id.get("vdi")
        vdi = current_obj_ts.variant_dependent_images[index]
        value_key = ctx.triggered_id.get("image")
        new_value = ctx.triggered[0]["value"]
        lb = vdi.images[value_key]["lower"]
        ub = vdi.images[value_key]["upper"]
        if ctx.triggered_id["type"] == "vdi_image_lb":
            lb = new_value
        else:
            ub = new_value
        vdi.set_image_bounds(value_key, lb, ub)
    elif isinstance(ctx.triggered_id, dict) and ctx.triggered_id.get("type") in ("vdt_table_lb", "vdt_table_ub"):
        index = ctx.triggered_id.get("vdt")
        vdt = current_obj_ts.variant_dependent_tables[index]
        value_key = ctx.triggered_id.get("table")
        new_value = ctx.triggered[0]["value"]
        lb = vdt.tables[value_key]["lower"]
        ub = vdt.tables[value_key]["upper"]
        if ctx.triggered_id["type"] == "vdt_table_lb":
            lb = new_value
        else:
            ub = new_value
        vdt.set_table_bounds(value_key, lb, ub)

    elif ctx.triggered_id == "i_title" and title:
        current_obj_ts.title = title
    elif ctx.triggered_id == "i_text" and text:
        item_body.update_item(current_obj, text)
    return item_body.items, test_structure.to_dict(), True