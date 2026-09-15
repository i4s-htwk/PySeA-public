from dash import Output, Input, State, callback, dcc, html, MATCH, ALL, ctx
from dash.exceptions import PreventUpdate
from openpyxl import load_workbook
import os

from backend.stores import StoreTables, StoreConfigurations
from backend.utils import CustomTable, ExcelTable


def render_excel_table(list_tables, show_tables, count_excel, i, path_excel_file, test_structure_obj):
    options_excel = options_page = []
    if path_excel_file:
        options_excel = [{"label": path.split("/")[-1], "value": path} for path in path_excel_file if path]
    if list_tables[i].page:
        options_page = [{"label": list_tables[i].page, "value": list_tables[i].page}]
    bereiche = []
    for j, area in enumerate(list_tables[i].list_areas):
        bereiche.append(
        html.Div([
            html.Div(area.id,style={"marginTop":"5px", "marginRight":"10px"}),html.Button("x", id={"type": "b_delete_area", "table":i, "index": j}),
            dcc.Input(id={"type": "von_zelle", "table":i, "count_excel":count_excel, "index":j}, type="text", value=area.cell_tl.value, placeholder="von Zelle", style={"marginLeft": "20px", "width": "200px", "fontSize": "14px"}),
            dcc.Input(id={"type": "bis_zelle", "table":i, "count_excel":count_excel, "index":j}, type="text", value=area.cell_br.value, placeholder="bis Zelle", style={"marginLeft": "10px", "marginRight": "25px", "width": "200px", "fontSize": "14px"}),
            show_excel_table(area, list_tables[i], show_tables)],
            style={"marginLeft": "80px", "marginTop":"10px","display": "flex", "alignItems": "flex-start"}))

    automatic_responses = dcc.Checklist(id={"type": "b_automatic_responses", "table":i, "count_excel":count_excel}, options=[{"label": "Lücken automatisch generieren", "value": "true"}], value=["true"] if list_tables[i].automatic_responses else [], style={"marginLeft": "20px"})
    if "section" in test_structure_obj:
        automatic_responses.style = {"display":"none"}


    return html.Div([
                html.Div([
                     html.H3("Exceltabelle", style={"margin":"0"}), html.Button("x", id={"type": "b_delete_excel_table", "index": i}, style={"marginLeft": "10px"}),
                     dcc.Dropdown(id={"type": "excel", "table":i, "count_excel":count_excel}, options=options_excel, value=list_tables[i].excel_file, placeholder="Exceldatei auswählen", style={"paddingLeft": "100px", "width": "200px", "height": "15px", "fontSize": "12px"}),
                     dcc.Dropdown(id={"type": "seite", "table":i, "count_excel":count_excel}, options=options_page, value=list_tables[i].page, placeholder="Seite auswählen", style={"paddingLeft": "20px", "width": "200px", "height": "15px", "fontSize": "12px"}),
                    automatic_responses],
                    style={"display": "flex",  "alignItems": "flex-start", "paddingBottom":"15px"}),
                html.Div([html.H4("Bereich", style={"margin": "20px 400px 0 30px","paddingLeft":"180px"}), html.H4("Vorschau", style={"margin": "20px 0 0 0"})], style={"display": "flex", "alignItems": "flex-start"}),
                html.Div(bereiche),
                html.Button("neuer Bereich", id={"type": "neuer_bereich", "table":i}, style={"width": "150px", "marginLeft": "210px", "marginTop": "10px"})
        ], style={"marginTop": "6px", "padding":"10px","border": "0.5px solid black"})

def render_custom_table(table, count_custom, i, test_structure_obj):
    automatic_responses = dcc.Checklist(id={"type": "b_automatic_responses", "table": i, "count_excel": count_custom},
                                        options=[{"label": "Lücken automatisch generieren", "value": "true"}],
                                        value=["true"] if table.automatic_responses else [],
                                        style={"marginLeft": "20px"})
    if "section" in test_structure_obj:
        automatic_responses.style = {"display": "none"}

    return html.Div([
        html.Div([
            html.Div([
                html.H3("Benutzerdef. Tabelle", style={"margin":"0"}),
                html.Button("x", id={"type": "b_delete_excel_table", "index": i}, style={"marginLeft": "10px"})],
            style={"display":"flex", "flexDirection":"row", "paddingBottom":"10px"}),
            table.id
        ], style={"marginRight": "20px"}),

        html.Div([
            html.H4("Reihenanzahl", style={"padding": 0, "margin": 0}),
            dcc.Input(id={"type":"i_row_count_custom_table", "table":i, "count_custom":count_custom}, value=len(table.cells), type="number", style={"width": "60px", "marginTop": "10px"})
        ], style={"display":"flex", "flexDirection":"column", "alignItems":"flex-start", "marginLeft":"15px"}),

        html.Div([
            html.H4("Spaltenanzahl", style={"padding": 0, "margin": 0}),
            dcc.Input(id={"type":"i_col_count_custom_table", "table":i, "count_custom":count_custom}, value=len(table.cells[0]) if table.cells else 0, type="number", style={"width": "60px", "marginTop": "10px"})
        ], style={"display":"flex", "flexDirection":"column", "alignItems":"flex-start", "marginLeft":"15px"}),
        automatic_responses,
        show_custom_table(table, i, count_custom)
    ], style={"marginTop": "6px", "padding": "10px", "border": "0.5px solid black", "display": "flex", "alignItems": "flex-start"})

@callback(
    Output("div_tables", "children"),
    Output("view_edit_item", "data", allow_duplicate=True),

    Input("b_tables", "n_clicks"),
    State("store_tables", "data"),
    State("store_configurations", "data"),
    State("store_current_obj", "data"),
    prevent_initial_call=True
)
def render_tables(n_clicks, store_tables, store_configurations, current_obj):
    if not ctx.triggered_id:
        raise PreventUpdate
    if not n_clicks:
        raise PreventUpdate

    tables = StoreTables.from_dict(store_tables or {})
    configurations = StoreConfigurations.from_dict(store_configurations or {})
    list_tables = tables.get_tables(current_obj)

    count_custom = count_excel = 0
    children_list = []
    for i, table in enumerate(list_tables):
        if isinstance(table, ExcelTable):
            children_list.append(render_excel_table(list_tables, tables.get_show_tables(current_obj), count_excel, i, configurations.path_excel_files, current_obj))
            count_excel += 1
        elif isinstance(table, CustomTable):
            children_list.append(render_custom_table(table, count_custom, i, current_obj))
            count_custom += 1

    return html.Div([
            html.Div(id="b_tablesliste", children=children_list),
            html.Div([
                html.Div([
                    dcc.Dropdown(id="dd_table_type", options=["Exceltabelle", "Benuterdefinierte Tabelle"], value=tables.tabletype, placeholder="Tabellenart wählen", style={"width":"200px", "height":"20px"}),
                    html.Button("neue Tabelle", id="b_new_table", style={"width": "200px", "padding": "0", "marginTop": "20px"})]
                ,style={"display":"flex", "flex-direction":"column", "paddingTop":"20px"}),
                dcc.Checklist(id="b_show_tables", options=[{"label": "Tabellenvorschau", "value": "show"}], value=["show"] if tables.get_show_tables(current_obj) else [], style={"marginLeft":"20px", "marginTop":"20px"})],
            style= {"display":"flex", "alignItems":"flex-start"})],
        style={"marginLeft": "260px", "marginTop": "60px", "paddingLeft": "20px"}),  "tables"

def show_excel_table(area, table, b_show_tables):
    if not b_show_tables:
        return html.Div()
    if not (area.cell_tl.value and area.cell_tl.is_excel_cell() and area.cell_br.value and area.cell_br.is_excel_cell() and table.page and table.excel_file):
        return html.Div("Vorschau konnte nicht geladen werden. Überprüfen Sie die Eingaben.", style={"marginTop":"5px"})
    table.create_excel_table(area.id)
    return html.Div([
        html.Table([
            html.Tbody([
                html.Tr([
                    html.Td(str(round(zelle.value, 5) if type(zelle.value)==float else zelle.value), rowSpan = zelle.rowspan, colSpan= zelle.colspan, style={"border": "1px solid black", "padding":"4px"}) for zelle in zeile
                ]) for zeile in area.cells
            ])], style={"borderCollapse": "collapse"})])

def show_custom_table(table, table_count, custom_count):
    return html.Table([
        html.Tbody([
            html.Tr([
                html.Td(
                    dcc.Input(
                        id={
                            "type": "i_cell_custom_table",
                            "table": table_count,
                            "count_custom": custom_count,
                            "row": i,
                            "col": j
                        },
                        value=table.cells[i][j].value,
                        style={
                            "width": "120px",
                            "minWidth": "120px",
                            "boxSizing": "border-box"
                        }
                    ),
                    style={
                        "border": "1px solid black",
                        "padding": "2px",
                        "minWidth": "120px"
                    }
                ) for j in range(len(table.cells[i]))
            ]) for i in range(len(table.cells))
        ])
    ], style={
        "borderCollapse": "collapse",
        "marginLeft": "50px",
        "tableLayout": "fixed"
    })

@callback(
    Output({"type": "seite", "table": MATCH, "count_excel":MATCH}, "options"),

    Input({"type": "excel", "table": MATCH, "count_excel":MATCH}, "value"),
    Input("tables_is_saved", "data"),
    prevent_initial_call=True
)
def lade_seitenoptionen(excel_path_file, is_saved):
    if not excel_path_file or not os.path.exists(excel_path_file):
        raise PreventUpdate
    try:
        wb = load_workbook(excel_path_file, read_only=True)
        seiten = wb.sheetnames
        wb.close()
        options = [{"label": name, "value": name} for name in seiten]
        return options
    except Exception as e:
        print("Fehler beim Einlesen der Excel-Datei:", e)
        return []

@callback(
    Output("tables_is_saved", "data", allow_duplicate=True),
    Output("div_tables", "children", allow_duplicate=True),

    Input("tables_is_saved", "data"),
    State("store_tables", "data"),
    State("store_configurations", "data"),
    State("store_current_obj", "data"),
    prevent_initial_call=True)
def rerender_tables(is_saved, store_tables, store_configurations, current_obj):
    if not ctx.triggered_id or not is_saved:
        raise PreventUpdate
    layout, *_ = render_tables(1, store_tables, store_configurations, current_obj)
    return False, layout

@callback(
    Output("store_tables", "data", allow_duplicate=True),
    Output("tables_is_saved", "data", allow_duplicate=True),

    Input({"type": "excel", "table": ALL, "count_excel":ALL}, "value"),
    Input({"type": "seite", "table": ALL, "count_excel":ALL}, "value"),
    Input({"type": "b_automatic_responses", "table": ALL, "count_excel":ALL}, "value"),
    Input({"type": "von_zelle", "table": ALL, "count_excel":ALL, "index": ALL}, "value"),
    Input({"type": "bis_zelle", "table": ALL, "count_excel":ALL, "index": ALL}, "value"),
    Input("b_show_tables", "value"),
    Input({"type": "neuer_bereich", "table": ALL}, "n_clicks"),
    Input("b_new_table", "n_clicks"),
    Input("dd_table_type", "value"),
    Input({"type": "b_delete_excel_table", "index": ALL}, "n_clicks"),
    Input({"type": "b_delete_area", "table":ALL, "index": ALL}, "n_clicks"),

    Input({"type":"i_row_count_custom_table", "table":ALL, "count_custom":ALL}, "value"),
    Input({"type":"i_col_count_custom_table", "table":ALL, "count_custom":ALL}, "value"),
    Input({"type": "i_cell_custom_table", "table": ALL, "count_custom":ALL, "row": ALL, "col": ALL}, "value"),

    State("store_tables", "data"),
    State("store_current_obj", "data"),
    prevent_initial_call=True
)
def speichere_b_tables(excel_files, pages, automatic_responses, cells_tl, cells_br, show_tables,
                       new_area, new_table, table_type, delete_excel_table, delete_area, row_counts, col_counts, cell_custom_table,
                       data, current_obj):
    if not ctx.triggered_id:
        raise PreventUpdate

    tables = StoreTables.from_dict(data or {})

    tables.set_show_tables(current_obj, "show" in show_tables)
    list_tables = tables.get_tables(current_obj)

    if isinstance(ctx.triggered_id, dict) and ctx.triggered_id.get("type") in ("excel", "seite"):
        table = ctx.triggered_id.get("table")
        count = ctx.triggered_id.get("count_excel")
        tab = list_tables[table]

        tab.excel_file = excel_files[count]
        tab.page = pages[count] if tab.excel_file else ""
        tab.automatic_responses = bool(automatic_responses[count])
    elif isinstance(ctx.triggered_id, dict) and ctx.triggered_id.get("type") == "b_automatic_responses":
        table = ctx.triggered_id.get("table")
        tab = list_tables[table]
        triggered_value = ctx.triggered[0].get("value", None)
        tab.automatic_responses = bool(triggered_value)

    elif isinstance(ctx.triggered_id, dict) and ctx.triggered_id.get("type") in ("von_zelle", "bis_zelle"):
        table_index = ctx.triggered_id.get("table")
        area_index = ctx.triggered_id.get("index")
        tab = list_tables[table_index]
        area = tab.list_areas[area_index]

        triggered_value = None
        if ctx.triggered:
            triggered_value = ctx.triggered[0].get("value", None)
        if triggered_value is None:
            triggered_value = ""

        if ctx.triggered_id.get("type") == "von_zelle":
            area.cell_tl.value = triggered_value
        else:
            area.cell_br.value = triggered_value

    elif isinstance(ctx.triggered_id, dict) and ctx.triggered_id.get("type") in ("i_row_count_custom_table", "i_col_count_custom_table"):
        table_index = ctx.triggered_id.get("table")
        count = ctx.triggered_id.get("count_custom")
        if row_counts[count]!="" and col_counts[count]!="":
            tab = list_tables[table_index]
            tab.update_table(int(row_counts[count]), int(col_counts[count]))

    elif isinstance(ctx.triggered_id, dict) and ctx.triggered_id.get("type") == "i_cell_custom_table":
        table_index = ctx.triggered_id.get("table")
        row = ctx.triggered_id.get("row")
        col = ctx.triggered_id.get("col")
        tab = list_tables[table_index]
        triggered_value = ctx.triggered[0].get("value", "") if ctx.triggered else ""
        if triggered_value is None:
            triggered_value = ""
        if 0 <= row < len(tab.cells) and 0 <= col < len(tab.cells[row]):
            tab.cells[row][col].value = triggered_value

    elif isinstance(ctx.triggered_id, dict) and ctx.triggered_id.get("type") == "neuer_bereich" and new_area:
        index = ctx.triggered_id.get("table")
        list_tables[index].add_area()

    elif ctx.triggered_id == "dd_table_type":
        tables.tabletype = table_type

    elif ctx.triggered_id == "b_new_table" and new_table:
        if table_type == "Exceltabelle":
            #list_tables.append("excelTable")
            tables.add_table(current_obj, "excelTable")
        elif table_type == "Benuterdefinierte Tabelle":
            #list_tables.append("customTable")
            tables.add_table(current_obj, "customTable")

    elif isinstance(ctx.triggered_id, dict) and ctx.triggered_id.get("type") == "b_delete_excel_table":
        index = ctx.triggered_id.get("index")
        if delete_excel_table and delete_excel_table[index] and delete_excel_table[index] > 0:
                list_tables.pop(index)

    elif isinstance(ctx.triggered_id, dict) and ctx.triggered_id.get("type") == "b_delete_area" and delete_area:
        table = ctx.triggered_id.get("table")
        index = ctx.triggered_id.get("index")
        if len(list_tables[table].list_areas) > 1:
            list_tables[table].list_areas.pop(index)

    list_tables = tables.get_tables(current_obj)
    for i, table in enumerate(list_tables):
        table.id = f"TABELLE_{i + 1}"
        if isinstance(table, ExcelTable):
            for j, area in enumerate(table.list_areas):
                area.id = table.id + f"_{j + 1}"

    return tables.to_dict(), True
