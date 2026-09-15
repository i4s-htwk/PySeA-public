import os
from dash import Output, Input, State, callback, dcc, html, ALL, ctx
from dash.exceptions import PreventUpdate
from openpyxl import load_workbook

from backend.stores import StoreResponses, StoreConfigurations

def render_responses(responses):
    return html.Div([
        html.Div(id="container_answer_list", children=[
            html.Div([
                html.Div(answer.id),
                dcc.Input(id={"type": "answer", "index": i}, type="text", value=answer.value,
                          placeholder="Lösung", style={"marginLeft": "20px", "width": "200px"}),
                html.Button("x", id={"type": "b_remove_answer", "index": i},
                            style={"marginLeft": "10px", "marginTop": "2px"})
            ], style={"display": "flex", "alignItems": "center", "marginBottom": "20px"})
            for i, answer in enumerate(responses)
        ]),
        html.Button("neue Antwort", id="b_new_response", style={"width": "250px", "padding": "0"})
    ], style={"display": "inline-block", "verticalAlign": "top", "marginRight": "20px", "border": "0.5px solid black", "padding":"10px", "width":"400px"})

def render_excel_responses(excel_responses, path_excel_file):
    if excel_responses:
        options_excel = options_page = []
        if path_excel_file:
            options_excel = [{"label": path.split("/")[-1], "value": path} for path in path_excel_file if path]
        if excel_responses.page:
            options_page = [{"label": excel_responses.page, "value": excel_responses.page}]

        div_excel_responses = html.Div([html.Div([
                        html.Div(resp.id),
                        dcc.Input(id={"type": "excel_answer", "index": i}, type="text", value=resp.cell, placeholder="Zelle in Exceldatei, z.B. 'B2'", style={"marginLeft": "20px", "width": "200px"}),
                        dcc.Input(id={"type": "excel_answer_points", "index": i}, type="text", value=resp.points, style={"marginLeft": "15px", "width": "20px"}),], style={"display": "flex", "alignItems": "center", "marginBottom":"20px"},
                ) for i, resp in enumerate(excel_responses.responses)])
        return html.Div([
                html.Div([
                    dcc.Dropdown(id = "excel", options=options_excel, value=excel_responses.excel_file, placeholder="Exceldatei auswählen", style={"width": "200px", "height": "15px", "fontSize": "12px"}),
                    dcc.Dropdown(id="page", options=options_page, value=excel_responses.page, placeholder="Seite auswählen", style={"paddingLeft": "20px", "width": "200px", "height": "15px", "fontSize": "12px"})],
                    style={"display":"flex", "flexDirection":"row", "marginBottom": "30px"}),
                dcc.Checklist(id="cb_acc_from_excel", options=[{"label": "Fehlertoleranz aus Exceldatei übernehmen", "value": "true"}], value=["true"] if excel_responses.answer_acc_from_excel else [], style={"marginBottom": "25px"}),
                div_excel_responses,
                html.Button("neue Antwort", id="b_new_excel_response", style={"display":"flex"}),
                html.Button(id="b_initialize_excel_responses", style={"display":"none"})],
                style={"display": "inline-block", "verticalAlign": "top" ,"border": "0.5px solid black", "padding":"10px", "width":"500px"})
    else:
        return html.Div([html.Button("Antworten aus Exceldatei", id="b_initialize_excel_responses"),
                         html.Button(id="b_new_excel_response", style={"display":"none"}), dcc.Dropdown(id="excel", style={"display":"none"}), dcc.Dropdown(id="page", style={"display":"none"}), dcc.Checklist(id="cb_acc_from_excel",style={"display":"none"})],
                        style={"display": "inline-block", "verticalAlign": "top", "padding":"10px"})

@callback(
    Output("div_answers", "children"),
    Output("view_edit_item", "data", allow_duplicate=True),

    Input("b_answers", "n_clicks"),
    State("store_responses", "data"),
    State("store_configurations", "data"),
    State("store_current_obj", "data"),
    prevent_initial_call=True
)
def render_answers(n_clicks, store_responses, store_configurations, current_obj):
    if not (ctx.triggered_id and n_clicks):
        raise PreventUpdate

    responses = StoreResponses.from_dict(store_responses or {})
    configurations = StoreConfigurations.from_dict(store_configurations or {})

    div_responses = render_responses(responses.items[current_obj].responses)
    div_excelresponses = render_excel_responses(responses.items[current_obj].excel_responses, configurations.path_excel_files)
    return html.Div([div_responses, div_excelresponses],style={"display":"block", "marginLeft": "260px", "marginTop": "80px", "paddingLeft": "20px"}), "answers"

@callback(
    Output("page", "options"),

    Input("excel", "value"),
    Input("answers_is_saved", "data"),
    prevent_initial_call=True
)
def lade_seitenoptionen(excel_path_file, is_saved):
    if not excel_path_file or not os.path.exists(excel_path_file):
        raise PreventUpdate
    try:
        wb = load_workbook(excel_path_file, read_only=True)
        seiten = wb.sheetnames
        wb.close()
        options = [{"label": name, "value": name} for name in seiten]
        return options
    except Exception as e:
        print("Fehler beim Einlesen der Excel-Datei:", e)
        return []

@callback(
    Output("answers_is_saved", "data", allow_duplicate=True),
    Output("div_answers", "children", allow_duplicate=True),

    Input("answers_is_saved", "data"),
    State("store_responses", "data"),
    State("store_configurations", "data"),
    State("store_current_obj", "data"),
    prevent_initial_call=True
)
def rerender_answers(is_saved, store_responses, store_configurations, current_obj):
    if not ctx.triggered_id or not is_saved:
        raise PreventUpdate
    layout, *_ = render_answers(1, store_responses, store_configurations, current_obj)
    return False, layout

@callback(
    Output("store_responses", "data", allow_duplicate=True),
    Output("answers_is_saved", "data", allow_duplicate=True),

    Input({"type": "answer", "index": ALL}, "value"),
    Input({"type": "excel_answer", "index": ALL}, "value"),
    Input({"type": "excel_answer_points", "index": ALL}, "value"),
    Input ("excel", "value"),
    Input ("page", "value"),
    Input("cb_acc_from_excel", "value"),
    Input({"type": "b_remove_answer", "index": ALL}, "n_clicks"),
    Input("b_new_response", "n_clicks"),
    Input("b_new_excel_response", "n_clicks"),
    Input("b_initialize_excel_responses", "n_clicks"),

    State("store_responses", "data"),
    State("store_current_obj", "data"),
    prevent_initial_call=True
)
def save_answers(responses, excel_responses, excel_points, excel, page, acc_from_excel, remove_answer, new_response, new_excel_response, init_excel_responses, store_responses_obj, current_obj):
    if not ctx.triggered_id:
        raise PreventUpdate

    store_responses_obj = StoreResponses.from_dict(store_responses_obj)
    responses_obj = store_responses_obj.items[current_obj]
    excel_responses_obj = responses_obj.excel_responses

    if isinstance(ctx.triggered_id, dict) and ctx.triggered_id.get("type") == "b_remove_answer":
        index = ctx.triggered_id.get("index")
        responses_obj.delete_response("RESPONSE_" + str(index + 1))
    elif isinstance(ctx.triggered_id, dict) and ctx.triggered_id.get("type") == "answer":
        index = ctx.triggered_id.get("index")
        responses_obj.set_response_value("RESPONSE_" + str(index + 1), responses[index])
    elif isinstance(ctx.triggered_id, dict) and ctx.triggered_id.get("type") == "excel_answer":
        index = ctx.triggered_id.get("index")
        excel_responses_obj.set_response_value("EXCEL_RESPONSE_" + str(index + 1), excel_responses[index])
    elif isinstance(ctx.triggered_id, dict) and ctx.triggered_id.get("type") == "excel_answer_points":
        index = ctx.triggered_id.get("index")
        points = excel_points[index] if excel_points[index]!="" else None
        excel_responses_obj.set_response_points("EXCEL_RESPONSE_" + str(index + 1), points)
    elif ctx.triggered_id == "excel" and excel:
        excel_responses_obj.excel_file = excel
    elif ctx.triggered_id == "page" and page:
        excel_responses_obj.page = page
    elif ctx.triggered_id == "cb_acc_from_excel":
        excel_responses_obj.answer_acc_from_excel = "true" in acc_from_excel
    elif ctx.triggered_id == "b_new_response" and new_response:
        responses_obj.add_response()
    elif ctx.triggered_id == "b_new_excel_response" and new_excel_response:
        excel_responses_obj.add_response()
    elif ctx.triggered_id == "b_initialize_excel_responses":
        responses_obj.add_excel_responses()

    return store_responses_obj.to_dict(), True

